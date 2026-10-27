
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from sqlalchemy import inspect
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Employee, Attendance, LeaveRequest, Payroll
from app.prediction_models import AttritionPrediction
from app.security_dependencies import require_roles

router = APIRouter(
    prefix="/api/reports",
    tags=["Reports and Exports"],
)

REPORT_MODELS = {
    "employees": Employee,
    "attendance": Attendance,
    "leave": LeaveRequest,
    "payroll": Payroll,
    "attrition": AttritionPrediction,
}

def normalize_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "value"):
        return value.value
    return value

def get_report_data(db: Session, report_type: str):
    model = REPORT_MODELS.get(report_type)

    if model is None:
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid report_type. Choose employees, attendance, "
                "leave, payroll, or attrition."
            ),
        )

    columns = [column.key for column in inspect(model).mapper.column_attrs]

    records = (
        db.query(model)
        .order_by(getattr(model, columns[0]).asc())
        .limit(10000)
        .all()
    )

    rows = []
    for record in records:
        rows.append([
            normalize_value(getattr(record, column))
            for column in columns
        ])

    return columns, rows

@router.get("/summary")
def report_summary(
    db: Session = Depends(get_db),
    current_user=Depends(require_roles("ADMIN", "HR")),
):
    summary = {}

    for report_type, model in REPORT_MODELS.items():
        summary[report_type] = db.query(model).count()

    return {
        "status": "success",
        "generated_at": datetime.utcnow().isoformat(),
        "report_counts": summary,
        "maximum_rows_per_export": 10000,
    }

@router.get("/export")
def export_report(
    report_type: str = Query(
        ...,
        description="employees, attendance, leave, payroll, or attrition",
    ),
    file_format: str = Query(
        "xlsx",
        description="xlsx or pdf",
    ),
    db: Session = Depends(get_db),
    current_user=Depends(require_roles("ADMIN", "HR")),
):
    report_type = report_type.strip().lower()
    file_format = file_format.strip().lower()

    if report_type not in REPORT_MODELS:
        raise HTTPException(
            status_code=400,
            detail="Unsupported report_type.",
        )

    if file_format not in ("xlsx", "pdf"):
        raise HTTPException(
            status_code=400,
            detail="file_format must be xlsx or pdf.",
        )

    columns, rows = get_report_data(db, report_type)
    output = BytesIO()

    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = report_type[:31].title()

        sheet.append(columns)

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

        for row in rows:
            sheet.append([
                "" if value is None else value
                for value in row
            ])

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            sheet.column_dimensions[letter].width = min(
                max(max_length + 2, 12), 30
            )

        workbook.save(output)
        media_type = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
        filename = f"{report_type}_report.xlsx"

    else:
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=20,
            leftMargin=20,
            topMargin=24,
            bottomMargin=24,
        )

        styles = getSampleStyleSheet()
        elements = [
            Paragraph(
                f"{report_type.title()} Report",
                styles["Title"],
            ),
            Paragraph(
                f"Generated: {datetime.utcnow().isoformat()} UTC",
                styles["Normal"],
            ),
            Paragraph(
                f"Records included: {len(rows)}",
                styles["Normal"],
            ),
            Spacer(1, 12),
        ]

        # Keep wide reports readable on landscape A4 pages.
        pdf_columns = columns[:8]
        pdf_rows = [
            [str(value or "")[:80] for value in row[:8]]
            for row in rows
        ]

        table = Table(
            [pdf_columns] + pdf_rows,
            repeatRows=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F2F6FA")]),
        ]))

        elements.append(table)
        document.build(elements)
        media_type = "application/pdf"
        filename = f"{report_type}_report.pdf"

    output.seek(0)

    return StreamingResponse(
        output,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )

